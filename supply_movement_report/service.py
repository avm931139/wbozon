from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import OPERATIONS_TG_BOT_TOKEN, OPERATIONS_TG_CHAT_ID, OPERATIONS_TG_PROXY_URL
from app.db import SessionLocal
from supply_movement_report.collectors import (
    ozon_returns, ozon_supplies, wb_returns, wb_supplies, yandex_movements,
)
from telegram_bot.client import TelegramClient


DEFAULT_PATH = Path("data/supply_movement_reports/marketplace_supplies_and_returns_all.xlsx")
MOSCOW = ZoneInfo("Europe/Moscow")


def _cell(value: Any) -> Any:
    if isinstance(value, datetime) and value.tzinfo:
        value = value.astimezone(MOSCOW).replace(tzinfo=None)
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value[:32760] if isinstance(value, str) else value


class SupplyMovementReportService:
    def __init__(self, *, session_factory: Callable[..., Any] = SessionLocal) -> None:
        self.session_factory = session_factory

    @staticmethod
    def _add_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
        ws = workbook.create_sheet(title)
        if not rows:
            ws.append(["Данных нет"])
            return
        columns = []
        for row in rows:
            for key in row:
                if key not in columns: columns.append(key)
        ws.append(columns)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for row in rows:
            ws.append([_cell(row.get(key)) for key in columns])
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for index, key in enumerate(columns, 1):
            width = 18 if len(key) < 18 else min(42, len(key) + 2)
            ws.column_dimensions[get_column_letter(index)].width = width

    def build(self, *, include_live_returns: bool = True) -> tuple[bytes, dict[str, Any]]:
        today = date.today(); errors = {}; data = {}
        with self.session_factory() as session:
            for key, callback in (("WB · Поставки", lambda: wb_supplies(session)), ("Ozon · Поставки", lambda: ozon_supplies(session))):
                try: data[key] = callback()
                except Exception as exc: errors[key] = f"{type(exc).__name__}: {exc}"; data[key] = []
            try:
                ym_supply, ym_return = yandex_movements(session, today)
                data["Яндекс · Поставки"] = ym_supply; data["Яндекс · Возвраты"] = ym_return
            except Exception as exc:
                errors["Яндекс · Заявки"] = f"{type(exc).__name__}: {exc}"
                data["Яндекс · Поставки"] = []; data["Яндекс · Возвраты"] = []
            if include_live_returns:
                for key, callback in (("WB · Возвраты", lambda: wb_returns(session, today)), ("Ozon · Возвраты", lambda: ozon_returns(session, today))):
                    try: data[key] = callback()
                    except Exception as exc: errors[key] = f"{type(exc).__name__}: {exc}"; data[key] = []
            else:
                data["WB · Возвраты"] = []; data["Ozon · Возвраты"] = []

        workbook = Workbook(); summary = workbook.active; summary.title = "Итог"
        summary.append(["Отчёт", "Поставки товаров и обратные движения за весь доступный период"])
        summary.append(["Сформирован", datetime.now(MOSCOW).replace(tzinfo=None)])
        summary.append(["Источник", "Поставки WB/Ozon — PostgreSQL; возвраты WB/Ozon и заявки Яндекса — актуальный API"])
        summary.append([]); summary.append(["Лист", "Строк", "Отправлено/план", "Принято/факт", "Статус источника"])
        for cell in summary[5]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
        order = ("WB · Поставки", "WB · Возвраты", "Ozon · Поставки", "Ozon · Возвраты", "Яндекс · Поставки", "Яндекс · Возвраты")
        for key in order:
            rows = data.get(key, [])
            planned = sum((row.get("Отправлено, шт.") or row.get("Отправлено/план, шт.") or row.get("К возврату, шт.") or row.get("Количество, шт.") or 0) for row in rows)
            actual = sum((row.get("Принято, шт.") or row.get("Принято/факт, шт.") or row.get("Фактически, шт.") or 0) for row in rows)
            summary.append([key, len(rows), planned, actual, errors.get(key, "OK")])
            self._add_sheet(workbook, key, rows)
        if errors:
            self._add_sheet(workbook, "Ошибки источников", [{"Источник": key, "Ошибка": value} for key, value in errors.items()])
        summary.column_dimensions["A"].width = 26; summary.column_dimensions["B"].width = 85
        stream = BytesIO(); workbook.save(stream); workbook.close()
        return stream.getvalue(), {"sheets": {key: len(data.get(key, [])) for key in order}, "errors": errors}

    def save(self, destination: str | Path = DEFAULT_PATH, *, include_live_returns: bool = True) -> dict[str, Any]:
        path = Path(destination).resolve(); path.parent.mkdir(parents=True, exist_ok=True)
        content, result = self.build(include_live_returns=include_live_returns)
        temporary = path.with_name(f".{path.name}.tmp")
        temporary.write_bytes(content); temporary.replace(path)
        return {"path": str(path), "bytes": len(content), **result}

    def send_telegram(self, destination: str | Path = DEFAULT_PATH) -> dict[str, Any]:
        result = self.save(destination)
        if not OPERATIONS_TG_BOT_TOKEN or not OPERATIONS_TG_CHAT_ID:
            raise RuntimeError("operations Telegram bot/chat is not configured")
        path = Path(result["path"])
        client = TelegramClient(OPERATIONS_TG_BOT_TOKEN, OPERATIONS_TG_CHAT_ID, proxy_url=OPERATIONS_TG_PROXY_URL)
        result["telegram_message_id"] = client.send_document(path.name, path.read_bytes(), caption="Поставки и возвраты по WB, Ozon и Яндекс Маркету за весь доступный период")
        return result
