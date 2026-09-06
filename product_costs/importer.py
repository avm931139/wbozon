from __future__ import annotations

import hashlib
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import MasterProduct, ProductCostImportRun, ProductCostRecord


SHEET_NAME = "Себестоимость"
EXPECTED_HEADERS = {
    1: "ID товара (не менять)",
    2: "Базовый артикул (не менять)",
    4: "Себестоимость 1 шт., ₽",
    5: "Количество на текущий момент, шт.",
    17: "Дата фиксации остатков",
}


class ProductCostImportError(ValueError):
    pass


class ProductCostImporter:
    def __init__(self, *, session_factory: Callable[..., Any] = SessionLocal) -> None:
        self.session_factory = session_factory

    def import_file(self, filename: str | Path) -> dict[str, Any]:
        path = Path(filename)
        payload = path.read_bytes()
        source_hash = hashlib.sha256(payload).hexdigest()
        with self.session_factory() as session:
            previous = session.query(ProductCostImportRun).filter_by(
                source_sha256=source_hash, status="completed"
            ).one_or_none()
            if previous:
                return {
                    "run_id": previous.id,
                    "status": "already_imported",
                    "rows_imported": previous.rows_imported,
                    "rows_skipped_blank": previous.rows_skipped_blank,
                }

        rows, total, skipped = self._read(path)
        run_id = uuid.uuid4().hex
        started = datetime.now(timezone.utc)
        with self.session_factory() as session:
            run = ProductCostImportRun(
                id=run_id,
                source_file=path.name,
                source_sha256=source_hash,
                started_at=started,
                status="running",
                rows_total=total,
                rows_skipped_blank=skipped,
            )
            session.add(run)
            try:
                products = {
                    row.id: row
                    for row in session.query(MasterProduct).filter(
                        MasterProduct.id.in_([item["master_product_id"] for item in rows]),
                        MasterProduct.active.is_(True),
                    )
                }
                for item in rows:
                    product = products.get(item["master_product_id"])
                    if product is None:
                        raise ProductCostImportError(
                            f"row {item['source_row']}: master product was not found"
                        )
                    if product.article != item["article"]:
                        raise ProductCostImportError(
                            f"row {item['source_row']}: article does not match product ID"
                        )
                    session.add(ProductCostRecord(import_run_id=run_id, created_at=started, **item))
                run.status = "completed"
                run.rows_imported = len(rows)
                run.finished_at = datetime.now(timezone.utc)
                session.commit()
            except Exception as exc:
                session.rollback()
                failed = ProductCostImportRun(
                    id=run_id,
                    source_file=path.name,
                    source_sha256=source_hash,
                    started_at=started,
                    finished_at=datetime.now(timezone.utc),
                    status="failed",
                    rows_total=total,
                    rows_imported=0,
                    rows_skipped_blank=skipped,
                    error=f"{type(exc).__name__}: {exc}",
                )
                session.add(failed)
                session.commit()
                raise
        return {
            "run_id": run_id,
            "status": "completed",
            "rows_imported": len(rows),
            "rows_skipped_blank": skipped,
        }

    @staticmethod
    def _read(path: Path) -> tuple[list[dict[str, Any]], int, int]:
        workbook = load_workbook(path, data_only=False, read_only=True)
        try:
            if SHEET_NAME not in workbook.sheetnames:
                raise ProductCostImportError(f"sheet {SHEET_NAME!r} was not found")
            sheet = workbook[SHEET_NAME]
            for column, expected in EXPECTED_HEADERS.items():
                if sheet.cell(1, column).value != expected:
                    raise ProductCostImportError(f"unexpected header in column {column}")
            rows = []
            skipped = 0
            seen: set[int] = set()
            for source_row in range(2, sheet.max_row + 1):
                raw_cost = sheet.cell(source_row, 4).value
                if raw_cost is None or (isinstance(raw_cost, str) and not raw_cost.strip()):
                    skipped += 1
                    continue
                master_product_id = ProductCostImporter._integer(
                    sheet.cell(source_row, 1).value, source_row, "product ID"
                )
                if master_product_id in seen:
                    raise ProductCostImportError(f"row {source_row}: duplicate product ID")
                seen.add(master_product_id)
                try:
                    unit_cost = Decimal(str(raw_cost)).quantize(
                        Decimal("0.000001"), rounding=ROUND_HALF_UP
                    )
                except (InvalidOperation, ValueError):
                    raise ProductCostImportError(f"row {source_row}: invalid unit cost")
                if unit_cost < 0:
                    raise ProductCostImportError(f"row {source_row}: unit cost is negative")
                quantity = ProductCostImporter._integer(
                    sheet.cell(source_row, 5).value, source_row, "quantity"
                )
                if quantity < 0:
                    raise ProductCostImportError(f"row {source_row}: quantity is negative")
                effective_at = sheet.cell(source_row, 17).value
                if not isinstance(effective_at, datetime):
                    raise ProductCostImportError(f"row {source_row}: invalid effective date")
                if effective_at.tzinfo is None:
                    from zoneinfo import ZoneInfo
                    effective_at = effective_at.replace(tzinfo=ZoneInfo("Europe/Moscow"))
                rows.append({
                    "master_product_id": master_product_id,
                    "article": str(sheet.cell(source_row, 2).value or "").strip(),
                    "product_name": sheet.cell(source_row, 3).value,
                    "unit_cost": unit_cost,
                    "quantity": quantity,
                    "currency": "RUB",
                    "effective_at": effective_at,
                    "source_row": source_row,
                    "note": sheet.cell(source_row, 18).value,
                })
            return rows, sheet.max_row - 1, skipped
        finally:
            workbook.close()

    @staticmethod
    def _integer(value: Any, row: int, label: str) -> int:
        if isinstance(value, bool):
            raise ProductCostImportError(f"row {row}: invalid {label}")
        try:
            result = int(value)
        except (TypeError, ValueError):
            raise ProductCostImportError(f"row {row}: invalid {label}")
        if Decimal(str(value)) != Decimal(result):
            raise ProductCostImportError(f"row {row}: {label} must be an integer")
        return result
