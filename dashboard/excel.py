from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


MARKET_NAMES = {"wb": "Wildberries", "ozon": "Ozon", "yandex_market": "Яндекс Маркет"}


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _finish(workbook: Workbook) -> bytes:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315B8A")
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[letter].width = width
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def operational_excel(data: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Показатели"
    sheet.append(["Площадка", "Заказы, шт.", "Заказы, ₽", "Отмены, шт.",
                  "Выкупы, шт.", "Выкупы, ₽", "Реклама, ₽", "Остаток, шт.",
                  "Остаток в себестоимости, ₽"])
    for key, values in data["marketplaces"].items():
        ads = data.get("ads", {}).get(key, {})
        stocks = data.get("stocks", {}).get(key, {})
        sheet.append([MARKET_NAMES[key], values.get("orders"), values.get("orders_amount"),
                      values.get("cancelled"), values.get("buyouts"), values.get("buyouts_amount"),
                      ads.get("spend"), stocks.get("units"), stocks.get("cost_value")])
    cabinet = workbook.create_sheet("Кабинетная аналитика")
    cabinet.append(["Площадка", "Покрытие полное", "Дней с данными", "Ожидалось дней",
                    "Заказано, шт.", "Заказано, ₽", "Выкуплено, шт.", "Выкуплено, ₽",
                    "Отменено, шт.", "Отменено, ₽", "Источник"])
    for key, values in data.get("cabinet_analytics", {}).items():
        cabinet.append([MARKET_NAMES[key], bool(values.get("complete")), values.get("coverage_days"),
                        values.get("expected_days"), values.get("ordered_items"), values.get("ordered_amount"),
                        values.get("purchased_items"), values.get("purchased_amount"),
                        values.get("cancelled_items"), values.get("cancelled_amount"), values.get("source")])
    series = workbook.create_sheet("Заказы по дням")
    series.append(["Дата", "Площадка", "Заказы, шт.", "Сумма, ₽"])
    for row in data.get("series", []):
        series.append([row.get("day"), MARKET_NAMES.get(row.get("marketplace"), row.get("marketplace")),
                       row.get("orders"), row.get("revenue")])
    return _finish(workbook)


def pnl_excel(data: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "P&L"
    summary.append(["Площадка", "Период с", "Период по", "Продажи, ₽", "Компенсации, ₽", "Выручка, ₽",
                    "Расходы МП, ₽", "К выплате, ₽", "Себестоимость, ₽", "Прибыль, ₽"])
    total = data["total"]
    period = data.get("period", {})
    summary.append(["ИТОГО", period.get("from"), period.get("to"), total.get("sales_revenue"), total.get("compensation"),
                    total.get("revenue"), total.get("expenses"), total.get("net_payout"),
                    total.get("cost_of_goods"), total.get("profit")])
    for key, values in data["marketplaces"].items():
        if not values.get("available"):
            summary.append([MARKET_NAMES[key], period.get("from"), period.get("to"), "Нет полного финансового отчёта"])
            continue
        summary.append([MARKET_NAMES[key], period.get("from"), period.get("to"), values.get("sales_revenue"), values.get("compensation"),
                        values.get("revenue"), values.get("expenses"), values.get("net_payout"),
                        values.get("cost_of_goods"), values.get("profit")])
    expenses = workbook.create_sheet("Расходы")
    expenses.append(["Площадка", "Статья", "Категория", "Сумма, ₽", "% выручки", "Источник"])
    for key, values in data["marketplaces"].items():
        for row in values.get("expense_lines", []):
            expenses.append([MARKET_NAMES[key], row.get("label"), row.get("category_label"),
                             row.get("amount"), row.get("share_percent"), row.get("source")])
    return _finish(workbook)


def stocks_excel(data: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Остатки"
    sheet.append(["Артикул", "Наименование", "Себестоимость, ₽", "Дата себестоимости",
                  "WB, шт.", "WB обновлено", "Ozon, шт.", "Ozon обновлено",
                  "Яндекс, шт.", "Яндекс обновлено"])
    for row in data.get("rows", []):
        sheet.append([_cell_value(value) for value in (
            row.get("article"), row.get("name"), row.get("unit_cost"),
            row.get("cost_updated_at"), row["wb"].get("quantity"),
            row["wb"].get("updated_at"), row["ozon"].get("quantity"),
            row["ozon"].get("updated_at"), row["yandex_market"].get("quantity"),
            row["yandex_market"].get("updated_at"),
        )])
    return _finish(workbook)
