from __future__ import annotations

import csv
import calendar
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


def normalize(value: Any) -> str:
    return re.sub(r"[^a-zа-яё0-9]+", "", str(value or "").casefold())


def money(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text[:19], fmt).date()
        except ValueError:
            pass
    return None


MARKETPLACE_MARKERS = {
    "wb": ("wildberries", "вайлдберриз", "wb", "rrdid", "номерстроки"),
    "ozon": ("ozon", "озон", "operationid", "типначисления"),
    "yandex_market": ("yandexmarket", "яндексмаркет", "transactionid", "shopsku", "вашsku"),
}

ALIASES = {
    "wb": {
        "id": ("rrdid", "rrd_id", "номерстроки", "номерстрокиотчета"),
        "amount": ("forpay", "кперечислениюпродавцуза реализованныйтовар", "кперечислениюпродавцу", "сумма"),
        "date": ("rrdate", "датаотчета", "датапродажи"),
        "type": ("selleropername", "обоснованиедляоплаты"),
    },
    "ozon": {
        "id": ("operationid", "operation_id", "номероперации", "идентификатороперации"),
        "amount": ("amount", "суммаоперации", "итоговаясуммаоперации", "начислено"),
        "date": ("accrualdate", "operationdate", "датаоперации", "дата начисления"),
        "type": ("accrualtype", "operationtype", "типначисления", "типоперации"),
    },
    "yandex_market": {
        "id": ("transactionid", "transaction_id", "идентификатортранзакции", "номерплатежа"),
        "amount": ("transactionsum", "transaction_sum", "сумматранзакции", "сумма"),
        "date": ("transactiondate", "transaction_date", "дататранзакции", "датаплатежа"),
        "type": ("transactiontype", "типтранзакции", "типплатежа"),
        "order": ("orderid", "order_id", "номерзаказа"),
        "sku": ("shopsku", "yoursku", "вашsku", "артикулмагазина"),
    },
}


@dataclass(frozen=True)
class ParsedRow:
    key: str
    amount: Decimal
    row_date: date | None
    sheet: str
    row_number: int
    raw: dict[str, Any]


@dataclass(frozen=True)
class ParsedExport:
    marketplace: str
    rows: list[ParsedRow]
    period_start: date | None
    period_end: date | None
    ignored_sheets: list[str]


def _field(row: dict[str, Any], marketplace: str, name: str) -> Any:
    indexed = {normalize(k): v for k, v in row.items()}
    for alias in ALIASES[marketplace].get(name, ()):
        if normalize(alias) in indexed:
            return indexed[normalize(alias)]
    return None


def _row_key(row: dict[str, Any], marketplace: str) -> str | None:
    identifier = _field(row, marketplace, "id")
    kind = _field(row, marketplace, "type")
    day = as_date(_field(row, marketplace, "date"))
    if identifier not in (None, ""):
        parts = [str(identifier).strip()]
        if marketplace == "ozon":
            parts.append(str(kind or "").strip())
        return "|".join(parts)
    if marketplace == "yandex_market":
        order = _field(row, marketplace, "order")
        sku = _field(row, marketplace, "sku")
        if order or sku:
            return "|".join((str(order or "").strip(), str(sku or "").strip(), str(kind or "").strip(), day.isoformat() if day else ""))
    return None


def detect_marketplace(path: Path, headers: Iterable[Any], forced: str | None = None) -> str:
    if forced:
        if forced not in ALIASES:
            raise ValueError(f"unsupported marketplace: {forced}")
        return forced
    haystack = normalize(path.name + " " + " ".join(str(x or "") for x in headers))
    scores = {name: sum(normalize(marker) in haystack for marker in markers) for name, markers in MARKETPLACE_MARKERS.items()}
    winner = max(scores, key=scores.get)
    if scores[winner] == 0 or list(scores.values()).count(scores[winner]) > 1:
        raise ValueError("marketplace is not recognized; prefix filename with wb_, ozon_ or yandex_market_")
    return winner


def _xlsx_tables(path: Path) -> tuple[list[tuple[str, list[dict[str, Any]]]], list[Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[tuple[str, list[dict[str, Any]]]] = []
    all_headers: list[Any] = []
    try:
        for sheet in workbook.worksheets:
            values = sheet.iter_rows(values_only=True)
            sample = []
            for _ in range(30):
                try:
                    sample.append(next(values))
                except StopIteration:
                    break
            best_index, best_score = None, 0
            for index, candidate in enumerate(sample):
                normalized = {normalize(x) for x in candidate if x not in (None, "")}
                score = sum(any(normalize(alias) in normalized for alias in aliases) for market in ALIASES.values() for aliases in market.values())
                if score > best_score:
                    best_index, best_score = index, score
            if best_index is None or best_score < 2:
                tables.append((sheet.title, []))
                continue
            headers = [str(x).strip() if x is not None else f"column_{i}" for i, x in enumerate(sample[best_index], 1)]
            all_headers.extend(headers)
            data = sample[best_index + 1 :] + list(values)
            rows = [dict(zip(headers, values)) for values in data if any(x not in (None, "") for x in values)]
            tables.append((sheet.title, rows))
    finally:
        workbook.close()
    return tables, all_headers


def _csv_tables(path: Path) -> tuple[list[tuple[str, list[dict[str, Any]]]], list[Any]]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise ValueError("CSV encoding is not supported")
    dialect = csv.Sniffer().sniff(text[:8192], delimiters=";,\t")
    rows = list(csv.DictReader(text.splitlines(), dialect=dialect))
    return [(path.stem, rows)], list(rows[0]) if rows else []


def parse_export(path: Path, forced_marketplace: str | None = None) -> ParsedExport:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        tables, headers = _xlsx_tables(path)
    elif suffix == ".csv":
        tables, headers = _csv_tables(path)
    else:
        raise ValueError("supported cabinet export formats: .xlsx and .csv")
    marketplace = detect_marketplace(path, headers, forced_marketplace)
    parsed: list[ParsedRow] = []
    ignored: list[str] = []
    for sheet, rows in tables:
        before = len(parsed)
        for row_number, row in enumerate(rows, 2):
            key = _row_key(row, marketplace)
            amount = money(_field(row, marketplace, "amount"))
            if key and amount is not None:
                parsed.append(ParsedRow(key, amount, as_date(_field(row, marketplace, "date")), sheet, row_number, row))
        if len(parsed) == before:
            ignored.append(sheet)
    if not parsed:
        raise ValueError(f"no comparable {marketplace} financial rows found")
    dates = [row.row_date for row in parsed if row.row_date]
    if not dates:
        for day, month, year in re.findall(r"(?<!\d)(\d{1,2})[-_.](\d{1,2})[-_.](20\d{2})(?!\d)", path.stem):
            try:
                dates.append(date(int(year), int(month), int(day)))
            except ValueError:
                pass
    if not dates:
        month_match = re.search(r"(?<!\d)(20\d{2})[-_.](0[1-9]|1[0-2])(?!\d)", path.stem)
        if month_match:
            year, month = map(int, month_match.groups())
            dates.extend((date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])))
    return ParsedExport(marketplace, parsed, min(dates) if dates else None, max(dates) if dates else None, ignored)
