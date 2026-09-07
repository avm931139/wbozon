from __future__ import annotations

import json
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from finance_reconciliation.parsing import ParsedExport, ParsedRow


@dataclass
class Comparison:
    matched: list[dict[str, Any]]
    missing_in_api: list[dict[str, Any]]
    missing_in_cabinet: list[dict[str, Any]]
    amount_mismatches: list[dict[str, Any]]
    cabinet_total: Decimal
    api_total: Decimal
    cabinet_row_count: int
    api_row_count: int

    @property
    def status(self) -> str:
        return "correct" if not (self.missing_in_api or self.missing_in_cabinet or self.amount_mismatches) else "mismatch"


def _aggregate(rows: list[ParsedRow]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        current = result.setdefault(row.key, {"key": row.key, "amount": Decimal(0), "count": 0, "locations": []})
        current["amount"] += row.amount
        current["count"] += 1
        current["locations"].append(f"{row.sheet}:{row.row_number}")
    return result


def compare(cabinet: list[ParsedRow], api: list[ParsedRow], tolerance: Decimal) -> Comparison:
    left, right = _aggregate(cabinet), _aggregate(api)
    matched, missing_api, missing_cabinet, mismatches = [], [], [], []
    for key in sorted(left.keys() | right.keys()):
        cabinet_row, api_row = left.get(key), right.get(key)
        if api_row is None:
            missing_api.append(cabinet_row)
        elif cabinet_row is None:
            missing_cabinet.append(api_row)
        else:
            difference = cabinet_row["amount"] - api_row["amount"]
            item = {
                "key": key,
                "cabinet_amount": cabinet_row["amount"],
                "api_amount": api_row["amount"],
                "difference": difference,
                "cabinet_count": cabinet_row["count"],
                "api_count": api_row["count"],
                "cabinet_locations": ", ".join(cabinet_row["locations"]),
                "api_locations": ", ".join(api_row["locations"]),
            }
            if abs(difference) > tolerance or cabinet_row["count"] != api_row["count"]:
                mismatches.append(item)
            else:
                matched.append(item)
    return Comparison(
        matched, missing_api, missing_cabinet, mismatches,
        sum((x.amount for x in cabinet), Decimal(0)),
        sum((x.amount for x in api), Decimal(0)),
        len(cabinet), len(api),
    )


def _safe(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and len(value) > 32760:
        return value[:32720] + " …[обрезано Excel]"
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = workbook.create_sheet(title[:31])
    if not rows:
        ws.append(["Нет строк"])
        return
    columns = list(rows[0])
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in rows:
        ws.append([_safe(row.get(column)) for column in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_report(path: Path, parsed: ParsedExport, comparison: Comparison, source_name: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Итог"
    rows = [
        ("Статус", "КОРРЕКТНО" if comparison.status == "correct" else "ЕСТЬ РАСХОЖДЕНИЯ"),
        ("Маркетплейс", parsed.marketplace),
        ("Исходный файл", source_name),
        ("Период", f"{parsed.period_start or 'не определён'} — {parsed.period_end or 'не определён'}"),
        ("Строк кабинета", comparison.cabinet_row_count),
        ("Строк API", comparison.api_row_count),
        ("Совпало", len(comparison.matched)),
        ("Нет в API", len(comparison.missing_in_api)),
        ("Нет в файле кабинета", len(comparison.missing_in_cabinet)),
        ("Не совпала сумма/кратность", len(comparison.amount_mismatches)),
        ("Сумма кабинета", comparison.cabinet_total),
        ("Сумма API", comparison.api_total),
        ("Разница", comparison.cabinet_total - comparison.api_total),
    ]
    for row in rows:
        summary.append(tuple(_safe(value) for value in row))
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 70
    _sheet(workbook, "Не найдено в API", comparison.missing_in_api)
    _sheet(workbook, "Нет в файле кабинета", comparison.missing_in_cabinet)
    _sheet(workbook, "Расхождения сумм", comparison.amount_mismatches)
    _sheet(workbook, "Совпавшие строки", comparison.matched)
    workbook.save(path)
    workbook.close()
