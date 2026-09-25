from datetime import datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models import MasterProduct, ProductCostRecord
from dashboard.__main__ import ABC_HTML, HTML, PNL_HTML, STOCKS_HTML
from dashboard.excel import abc_excel, operational_excel, pnl_excel, stocks_excel
from dashboard.service import DashboardService

def test_dashboard_period_defaults_to_current_day():
    start,end=DashboardService.period(None,None)
    assert start == end


def test_previous_period_has_equal_length_and_precedes_current():
    assert DashboardService.previous_period(
        DashboardService.period("2026-09-05", "2026-09-07")[0],
        DashboardService.period("2026-09-05", "2026-09-07")[1],
    ) == (
        DashboardService.period("2026-09-02", "2026-09-04")[0],
        DashboardService.period("2026-09-02", "2026-09-04")[1],
    )

def test_dashboard_rejects_reverse_or_oversized_period():
    with pytest.raises(ValueError): DashboardService.period("2026-09-02","2026-09-01")
    with pytest.raises(ValueError): DashboardService.period("2020-01-01","2026-09-01")
    with pytest.raises(ValueError): DashboardService.period("2099-01-01", "2099-01-01")


def test_operational_dashboard_contains_only_operational_metrics():
    assert "Выкупленные товары" in HTML
    assert "Финансовый результат" not in HTML
    assert "value.buyouts_amount" in HTML
    assert "Оперативный дашборд" in HTML
    assert "每日各平台订单额" in HTML
    assert "bar.empty" in HTML
    assert "value.amount" in HTML
    assert ".profit" not in HTML
    assert 'id="month" type="month"' in HTML
    assert "function selectMonth()" in HTML
    assert "start.setDate(now.getDate()-6)" in HTML
    assert "lastDay>now?now:lastDay" in HTML
    assert "monthInput.max=today.slice(0,7)" in HTML


def test_operational_dashboard_has_readable_semantic_comparisons_and_help():
    assert "Главное за период" in HTML
    assert "По площадкам" in HTML
    assert "сравнение с равным периодом" in HTML
    assert "points=false,neutral=false,label=''" in HTML
    assert "lowerBetter:true,points:true" in HTML
    assert "neutral:true,label:'расход'" in HTML
    assert 'data-tooltip=' in HTML
    assert "Доля = отменённые товары ÷ заказанные товары × 100%" in HTML
    assert "ROAS = атрибутированная сумма заказов ÷ рекламный расход" in HTML
    assert "Сегодняшний день ещё не завершён" in HTML
    assert "Доля общих заказов" in HTML
    assert "Средняя цена товара" in HTML
    assert "Среднее в день" in HTML
    assert "Прогноз заказов месяца" in HTML
    assert "function monthProjection" in HTML
    assert ".bar.wb" in HTML
    assert "marketKeys=['wb','ozon','yandex_market']" in HTML
    assert "height:${height}%" in HTML
    assert "WB Order Feed и FBO Orders" in HTML
    assert "Ozon Postings" in HTML
    assert "Yandex Market Orders API" in HTML
    assert "Оперативные события / 实时事件 ${info(" in HTML
    assert "Кабинетная аналитика / 后台分析 ${info(" in HTML
    assert "Данные поступают почти в реальном времени" in HTML
    assert "Это воронка заказов, созданных в выбранном периоде" in HTML
    assert "Поздние выкупы, доставки, отмены и возвраты пересчитываются задним числом" in HTML
    assert "Процент выкупа / отмен" in HTML
    assert "cabinetPurchaseRate=ratio(cabinetPurchased,cabinetOrdered)" in HTML
    assert "useCabinetHeadline=key==='wb'&&cabinet.complete" in HTML
    assert "headlineAmount/days" in HTML
    assert "кабинетная аналитика':'оперативные данные'" in HTML
    assert 'class="analytics-note"' not in HTML


def test_operational_queries_compare_product_units_with_product_units():
    import inspect

    source = inspect.getsource(DashboardService._operational_period_metrics)
    assert "order_items" in source
    assert "cancelled_items" in source
    assert 'values.get("cancelled_items")' in source
    assert 'values.get("order_items")' in source


def test_cabinet_analytics_is_separate_and_exposes_ozon_realization_price():
    import inspect

    source = inspect.getsource(DashboardService._cabinet_analytics)
    assert "wb_sales_funnel_daily" in source
    assert "wb_sales_funnel_period_products" in source
    assert "period_metrics.ordered_items ELSE product_metrics.product_ordered_items" in source
    assert "yandex_market_sales_analytics_daily" in source
    assert "seller_price*coalesce(p.quantity,0)" in source
    assert "average_realized_price" in source
    summary = inspect.getsource(DashboardService.summary)
    assert '"cabinet_analytics":cabinet' in summary
    assert "Кабинетная аналитика" in HTML
    assert "Оперативные события" in HTML
    assert "cabinetComparable=cabinet.complete&&oldCabinet.complete" in HTML
    assert "Реализовано по начислениям" in HTML


def test_pnl_dashboard_uses_only_financial_labels_and_separate_endpoint():
    assert "/api/pnl" in PNL_HTML
    assert "Точные финансовые отчёты" in PNL_HTML
    assert "Все расходы МП" in PNL_HTML
    assert "value.profit" in PNL_HTML
    assert "value.revenue" in PNL_HTML
    assert "ROAS" not in PNL_HTML
    assert "value.orders_amount" not in PNL_HTML


def test_pnl_defaults_to_month_through_latest_elapsed_sunday():
    assert "function defaultPnlPeriod(now)" in PNL_HTML
    assert "daysAfterSunday=last.getDay()||7" in PNL_HTML
    assert "new Date(last.getFullYear(),last.getMonth(),1)" in PNL_HTML
    assert "defaultPeriod=defaultPnlPeriod(now)" in PNL_HTML
    assert "now.getMonth()-1" not in PNL_HTML


def test_pnl_shows_expense_breakdown_comparisons_and_source_tooltips():
    import inspect

    assert 'class="info"' in PNL_HTML
    assert "Расшифровка расходов" in PNL_HTML
    assert "Расходы по категориям" in PNL_HTML
    assert "share_percent" in PNL_HTML
    assert "previous_total" in PNL_HTML
    assert "source_endpoint" in PNL_HTML
    assert "updated_at" in PNL_HTML
    assert "Процент = статья расходов ÷ выручка × 100%" in PNL_HTML
    assert 'class="total-layout"' in PNL_HTML
    assert 'class="total-summary"' in PNL_HTML
    assert 'class="total-breakdown"' in PNL_HTML
    assert "@media(max-width:1000px)" in PNL_HTML
    source = inspect.getsource(DashboardService._pnl_expense_breakdowns)
    assert "wb_financial_sales_rows" in source
    assert "ozon_finance_accruals" in source
    assert "jsonb_path_query" in source
    assert "ozon_finance_posting_accruals" in source
    assert "ozon_finance_accrual_types" in source
    assert "t.name='SaleCommission'" in source
    assert "yandex_market_finance_transactions" in source


def test_pnl_expense_lines_reconcile_to_authoritative_finance_total():
    lines = DashboardService._expense_lines(
        [{"key": "delivery", "label": "Доставка", "amount": 20}],
        revenue=200,
        expected_total=50,
        source="finance table",
    )

    assert sum(line["amount"] for line in lines) == 50
    assert lines[0]["category"] == "logistics"
    assert lines[0]["share_percent"] == 10
    assert lines[1]["key"] == "reconciliation_adjustment"


def test_pnl_expense_category_uses_stable_marketplace_operation_code():
    lines = DashboardService._expense_lines(
        [{"key": "PayPerClick", "label": "Оплата за клик", "amount": 25}],
        revenue=100,
        expected_total=25,
        source="Ozon finance detail",
    )

    assert lines[0]["category"] == "advertising"


def test_pnl_yandex_contra_expense_is_preserved_without_reconciliation_residual():
    lines = DashboardService._expense_lines(
        [
            {"key": "placement", "label": "Размещение товарных предложений", "amount": 100},
            {"key": "reversal", "label": "Возврат списания · Скидка за лояльность", "amount": -10},
        ],
        revenue=200,
        expected_total=90,
        source="Yandex united netting",
    )

    assert sum(line["amount"] for line in lines) == 90
    assert lines[0]["category"] == "commission"
    assert lines[1]["category"] == "discounts"
    assert all(line["key"] != "reconciliation_adjustment" for line in lines)


def test_pnl_yandex_promotion_discount_is_not_logistics():
    category, _ = DashboardService._expense_category(
        "Доставка (средняя миля) · Скидка за участие в совместных акциях"
    )

    assert category == "discounts"


def test_pnl_explicit_advertising_takes_priority_over_discount_wording():
    category, _ = DashboardService._expense_category(
        "Буст продаж · рекламное продвижение · скидка"
    )

    assert category == "advertising"


def test_stock_dashboard_has_three_market_images_and_refresh_dates():
    assert "/api/stocks?date=" in STOCKS_HTML
    assert "Фото WB" in STOCKS_HTML
    assert "Фото Ozon" in STOCKS_HTML
    assert "Фото Яндекс" in STOCKS_HTML
    assert "WB · остаток / обновлено" in STOCKS_HTML
    assert "Ozon · остаток / обновлено" in STOCKS_HTML
    assert "Яндекс · остаток / обновлено" in STOCKS_HTML
    assert "только товары с остатком" in STOCKS_HTML
    assert 'href="/stocks"' in HTML
    assert 'href="/stocks"' in PNL_HTML
    assert "Себестоимость, ₽" in STOCKS_HTML
    assert "/api/product-cost" in STOCKS_HTML
    assert "saveCost(masterId)" in STOCKS_HTML


def test_all_dashboard_pages_export_the_selected_filter_to_excel():
    assert "Выгрузить в Excel" in HTML
    assert "report:'operational',from:fromInput.value,to:toInput.value" in HTML
    assert "Выгрузить в Excel" in PNL_HTML
    assert "report:'pnl',from:from.value,to:to.value" in PNL_HTML
    assert "Выгрузить в Excel" in STOCKS_HTML
    assert "report:'stocks',date:dateInput.value" in STOCKS_HTML
    assert "Выгрузить в Excel" in ABC_HTML
    assert "report:'abc',from:from.value,to:to.value" in ABC_HTML


def test_dashboard_excel_exports_are_valid_workbooks():
    markets = {
        key: {"orders": 1, "orders_amount": 100, "cancelled": 0,
              "buyouts": 1, "buyouts_amount": 90}
        for key in ("wb", "ozon", "yandex_market")
    }
    operations = operational_excel({
        "marketplaces": markets,
        "ads": {key: {"spend": 5} for key in markets},
        "stocks": {key: {"units": 2, "cost_value": 50} for key in markets},
        "cabinet_analytics": {key: {"complete": True, "coverage_days": 1,
                                     "expected_days": 1, "ordered_items": 1,
                                     "ordered_amount": 100, "source": "API"} for key in markets},
        "series": [{"day": "2026-09-01", "marketplace": "wb", "orders": 1, "revenue": 100}],
    })
    operational_book = load_workbook(BytesIO(operations), read_only=True)
    assert operational_book.sheetnames == ["Показатели", "Кабинетная аналитика", "Заказы по дням"]
    assert operational_book["Показатели"]["A2"].value == "Wildberries"

    available = {"available": True, "sales_revenue": 100, "compensation": 1,
                 "revenue": 101, "expenses": 20, "net_payout": 81,
                 "cost_of_goods": 30, "profit": 51,
                 "expense_lines": [{"label": "Комиссия", "category_label": "Комиссия",
                                    "amount": 20, "share_percent": 19.8, "source": "API"}]}
    finance = pnl_excel({"period": {"from": "2026-09-01", "to": "2026-09-07"},
                         "total": available,
                         "marketplaces": {key: available for key in markets}})
    finance_book = load_workbook(BytesIO(finance), read_only=True)
    assert finance_book.sheetnames == ["P&L", "Расходы"]
    assert finance_book["P&L"]["B2"].value == "2026-09-01"

    aware_time = datetime(2026, 9, 1, 1, tzinfo=timezone.utc)
    stock_row = {"article": "SKU-1", "name": "Товар", "unit_cost": 30,
                 "cost_updated_at": aware_time}
    stock_row.update({key: {"quantity": 2, "updated_at": aware_time}
                      for key in markets})
    stock = stocks_excel({"rows": [stock_row]})
    stock_book = load_workbook(BytesIO(stock), read_only=True)
    assert stock_book.sheetnames == ["Остатки"]
    assert stock_book["Остатки"]["A2"].value == "SKU-1"
    assert stock_book["Остатки"]["D2"].value == "2026-09-01T01:00:00+00:00"


def test_abc_dashboard_uses_only_normalized_product_economics_layer():
    import inspect

    source = inspect.getsource(DashboardService.abc)
    assert "fact_product_economics_daily" in source
    assert "fact_product_economics_controls" in source
    assert "wb_financial_sales_rows" not in source
    assert "ozon_finance_accruals" not in source
    assert "yandex_market_finance_transactions" not in source
    assert "ABC по товарам" in ABC_HTML
    assert "Выручка" in ABC_HTML
    assert "Прибыль" in ABC_HTML
    assert "Реклама" in ABC_HTML
    assert "Логистика" in ABC_HTML
    assert "последний" not in ABC_HTML.lower()
    assert "function defaultFinancialPeriod(now)" in ABC_HTML
    assert "daysAfterSunday=last.getDay()||7" in ABC_HTML
    assert "defaultPeriod=defaultFinancialPeriod(now)" in ABC_HTML
    assert "new Date(now.getFullYear(),now.getMonth(),0)" not in ABC_HTML


def test_abc_categories_use_80_15_5_and_separate_losses():
    values = {"leader": 8000, "middle": 1500, "tail": 500, "zero": 0}
    assert DashboardService._abc_categories(values) == {
        "leader": "A", "middle": "B", "tail": "C", "zero": "—",
    }
    assert DashboardService._abc_categories(
        {"profit": 100, "loss": -1, "zero": 0}, loss_class=True
    ) == {"profit": "A", "loss": "У", "zero": "У"}


def test_abc_excel_contains_matrix_calculation_control_and_methodology():
    metric = {
        "available": True, "units": 2, "revenue_kopecks": 10000,
        "revenue_category": "A", "revenue_share_percent": 100,
        "profit_kopecks": 3000, "profit_category": "A",
        "profit_margin_percent": 30, "advertising_kopecks": 500,
        "advertising_drr_percent": 5, "logistics_kopecks": 1000,
        "logistics_share_percent": 10, "logistics_per_unit_kopecks": 500,
        "expense_kopecks": 4000, "cost_kopecks": 3000,
        "expense_allocation_method": "allocated_by_revenue",
        "advertising_allocation_method": "direct_product_report",
        "missing_cost_rows": 0,
    }
    payload = {
        "period": {"from": "2026-08-01", "to": "2026-08-31"},
        "rows": [{"article": "SKU-1", "name": "Товар", "is_unallocated": False,
                  "total": metric, "wb": metric, "ozon": None, "yandex_market": None}],
        "controls": {"wb": {"available": True, "coverage_days": 31, "expected_days": 31,
                             "revenue_kopecks": 10000, "revenue_actual_kopecks": 10000,
                             "revenue_delta_kopecks": 0, "profit_kopecks": 3000,
                             "profit_actual_kopecks": 3000, "profit_delta_kopecks": 0},
                     "ozon": {}, "yandex_market": {}},
        "methodology": {"source": "fact_product_economics_daily"},
    }
    book = load_workbook(BytesIO(abc_excel(payload)), read_only=True)
    assert book.sheetnames == ["ABC-матрица", "Расчёт по SKU", "Контроль", "Методика"]
    assert book["ABC-матрица"]["A2"].value == "SKU-1"
    assert book["ABC-матрица"]["Z2"].value is None
    assert book["Контроль"]["A2"].value == "Wildberries"


def test_stock_dashboard_cost_edit_appends_history_record():
    engine = create_engine("sqlite+pysqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, future=True)
    now = datetime.now(timezone.utc)
    with factory() as session:
        product = MasterProduct(
            article="SKU-1", name="Product", active=True,
            created_at=now, updated_at=now,
        )
        session.add(product)
        session.commit()
        product_id = product.id

    service = DashboardService(session_factory=factory)
    first = service.update_product_cost(product_id, "123.456")
    second = service.update_product_cost(product_id, "150.25")

    assert first["unit_cost"] == 123.456
    assert second["unit_cost"] == 150.25
    with factory() as session:
        rows = session.query(ProductCostRecord).order_by(ProductCostRecord.id).all()
        assert len(rows) == 2
        assert float(rows[-1].unit_cost) == 150.25
        assert rows[-1].note == "Изменено вручную на странице остатков"


def test_stock_details_uses_current_rows_and_historical_snapshots():
    import inspect

    source = inspect.getsource(DashboardService.stock_details)
    assert "wb_warehouse_remains" in source
    assert "wb_warehouse_remain_snapshots" in source
    assert "Всего находится на складах" in source
    assert "wb_fbs_stocks" not in source
    assert "wb_fbo_stocks" not in source
    assert "ozon_stocks" in source
    assert "yandex_market_stocks" in source
    assert "ozon_stock_snapshots" in source
    assert "yandex_market_stock_snapshots" in source
    assert "snapshot_date<=:d" in source
    assert "marketplace_product_media" in source


def test_wb_dashboard_stock_total_excludes_seller_fbs_stock():
    import inspect

    current = inspect.getsource(DashboardService._stocks)
    historical = inspect.getsource(DashboardService._historical_stocks)
    assert "wb_warehouse_remains" in current
    assert "wb_warehouse_remain_snapshots" in historical
    assert "Всего находится на складах" in current
    assert "Всего находится на складах" in historical
    assert "wb_fbs_stocks" not in current
    assert "wb_fbo_stocks" not in current
    assert "wb_fbs_stock_snapshots" not in historical
    assert "wb_fbo_stock_snapshots" not in historical


def test_operational_service_does_not_query_finance_ledgers():
    import inspect

    source = inspect.getsource(DashboardService._operational_period_metrics)
    assert "wb_financial_sales_rows" not in source
    assert "ozon_finance_accruals" not in source
    assert "yandex_market_finance_transactions" not in source
    assert 'values["data_status"] = "operational"' in source


def test_yandex_advertising_uses_actual_marketing_charges_and_checks_coverage():
    import inspect

    source = inspect.getsource(DashboardService._advertising_metrics)
    assert "yandex_market_finance_transactions" in source
    assert "product_or_service ILIKE" in source
    assert "marketing_finance" in source
    assert "jsonb_path_query" in source
    assert "exists(@.type_id)" in source
    assert "t.name IN ('PayPerClick','Promotion')" in source
    assert "coverage_days" in source
    assert "attribution_complete" in source
    assert "attributionComplete" in HTML
    assert "рекламные отчёты загружаются" in HTML
    assert "Списано по финансовым данным" in HTML
    assert "sources=4" in source


def test_dashboard_uses_historical_marketplace_order_sources():
    import inspect

    source = inspect.getsource(DashboardService._period_metrics)
    assert "wb_fbo_orders" in source
    assert "finance_buyouts" in source
    assert "OzonPosting" not in source
    assert "coalesce(price_with_discount,finished_price,total_price,0)" in source


def test_pnl_uses_finance_rows_for_yandex_cost_of_goods():
    import inspect

    source = inspect.getsource(DashboardService._period_metrics)
    assert "FROM yandex_market_finance_transactions" in source
    assert "yandex_market_order_items" not in source
    assert "united_netting" in source
    assert "product_events" in source
    assert "max(quantity) quantity" in source
    assert "NOT product_transaction" in source


def test_dashboard_uses_wb_finance_rows_for_closed_period_metrics():
    import inspect

    source = inspect.getsource(DashboardService._period_metrics)
    assert "retail_price_with_discount*quantity" in source
    assert "finance_buyouts_amount+compensation-net_payout expenses" in source
    assert "net_payout+advertising.amount" not in source
    assert "wb_finance_exact" in source
    assert "wb_sales_funnel_daily" in source
    assert 'wb["data_status"] = "preliminary"' in source
    assert "Бронирование товара через самовывоз" in source
    net_payout_sql = source.split("net_payout", 1)[0].rsplit("coalesce(sum(", 1)[-1]
    assert "deduction" in net_payout_sql
    assert "rebill_logistic_cost" not in net_payout_sql


def test_pnl_costs_follow_the_cost_effective_on_each_business_date():
    import inspect

    source = inspect.getsource(DashboardService._period_metrics)
    assert "LEFT JOIN LATERAL" in source
    assert "cost.effective_at::date<=s.rr_date::date" in source
    assert "cost.effective_at::date<=s.accrual_date" in source
    assert "cost.effective_at::date<=s.transaction_date" in source


def test_wb_pnl_expands_commission_and_does_not_duplicate_rebill_logistics():
    import inspect

    source = inspect.getsource(DashboardService._pnl_expense_breakdowns)
    assert "retail_price_with_discount*quantity-for_pay-acquiring_fee" in source
    assert '"key": "commission"' in source
    assert "coalesce(sum(delivery_service),0) logistics" in source
    assert "delivery_service+rebill_logistic_cost" not in source
    assert "paymentSchedule" in source
    assert '"key": "payment_schedule"' in source


def test_dashboard_uses_ozon_finance_accruals_for_buyouts_and_profit():
    import inspect

    source = inspect.getsource(DashboardService._period_metrics)
    assert "t.name='SaleCommission'" in source
    assert "finance_buyouts_amount" in source
    assert "sales.finance_buyouts_amount+ledger.compensation-ledger.net_accrual expenses" in source
    assert "p.seller_price<0" in source


def test_dashboard_labels_ozon_advertising_without_calling_it_revenue():
    assert "Расход рекламного API" in HTML
    assert "Атрибутированная сумма заказов" in HTML


def test_dashboard_uses_cabinet_wide_wb_and_ozon_order_analytics():
    import inspect

    source = inspect.getsource(DashboardService._cabinet_analytics)
    assert "wb_sales_funnel_account_daily" in source
    assert "item_delta" in source
    assert "FROM ozon_daily_sales" in source
    assert "ordered_units" in source
    assert "Источник кабинетного периода" in HTML
    assert "Отменено и возвращено по воронке" in HTML
    assert "Заказано по аналитике" in HTML
