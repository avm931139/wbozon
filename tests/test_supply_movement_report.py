from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from supply_movement_report.service import SupplyMovementReportService


class FakeSession:
    def __enter__(self): return self
    def __exit__(self, *_args): return None


def test_report_has_separate_marketplace_supply_and_return_sheets():
    supply = [{"Поставка": "S-1", "Артикул": "A", "Наименование": "Товар", "Отправлено, шт.": 10, "Принято, шт.": 9}]
    reverse = [{"Возврат": "R-1", "Артикул": "A", "Наименование": "Товар", "К возврату, шт.": 1, "Фактически, шт.": 1}]
    service = SupplyMovementReportService(session_factory=FakeSession)
    with (
        patch("supply_movement_report.service.wb_supplies", return_value=supply),
        patch("supply_movement_report.service.wb_returns", return_value=reverse),
        patch("supply_movement_report.service.ozon_supplies", return_value=supply),
        patch("supply_movement_report.service.ozon_returns", return_value=reverse),
        patch("supply_movement_report.service.yandex_movements", return_value=(supply, reverse)),
    ):
        content, result = service.build()

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Итог", "WB · Поставки", "WB · Возвраты", "Ozon · Поставки",
        "Ozon · Возвраты", "Яндекс · Поставки", "Яндекс · Возвраты",
    ]
    assert result["errors"] == {}
    assert workbook["WB · Поставки"]["C2"].value == "Товар"
    workbook.close()


def test_source_failure_does_not_block_other_marketplaces():
    service = SupplyMovementReportService(session_factory=FakeSession)
    with (
        patch("supply_movement_report.service.wb_supplies", side_effect=RuntimeError("WB unavailable")),
        patch("supply_movement_report.service.wb_returns", return_value=[]),
        patch("supply_movement_report.service.ozon_supplies", return_value=[]),
        patch("supply_movement_report.service.ozon_returns", return_value=[]),
        patch("supply_movement_report.service.yandex_movements", return_value=([], [])),
    ):
        content, result = service.build()

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    assert "Ошибки источников" in workbook.sheetnames
    assert "WB · Поставки" in result["errors"]
    workbook.close()
